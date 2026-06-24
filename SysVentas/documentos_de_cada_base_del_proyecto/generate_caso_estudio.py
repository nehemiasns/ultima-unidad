import sys
# pyrefly: ignore [missing-import]
from docx import Document
# pyrefly: ignore [missing-import]
from docx.shared import Pt, Inches

def crear_caso_estudio():
    doc = Document()
    
    # Title
    doc.add_heading('CASO DE ESTUDIO: SISTEMA DE VENTAS "SysVentas"', 0)
    
    # Integrantes
    doc.add_paragraph('INTEGRANTES:\nMaikol Nehemias Quispe Ramos\nYerson Michel Mamani Challco', style='Intense Quote')
    
    # Secciones
    doc.add_heading('1. INTRODUCCIÓN', level=1)
    doc.add_paragraph('La gestión eficiente de las ventas y el control de inventario son procesos fundamentales para la supervivencia y crecimiento de cualquier negocio comercial moderno. En muchos establecimientos minoristas, estos procesos aún se realizan de forma manual o con herramientas genéricas, lo que genera pérdida de información, errores en la facturación y desconocimiento real del stock disponible.')
    doc.add_paragraph('Para solucionar esta problemática, surge el proyecto SysVentas, un software de Punto de Venta (POS) diseñado específicamente para automatizar, agilizar y asegurar las transacciones comerciales del día a día, permitiendo a los administradores tomar mejores decisiones basadas en datos reales.')
    
    doc.add_heading('2. DESCRIPCIÓN DEL PROBLEMA', level=1)
    doc.add_paragraph('En las tiendas comerciales tradicionales, los empleados enfrentan diversos retos operativos:\n- Demora en la atención al cliente al buscar los precios manualmente.\n- Errores de cálculo al sumar el total de la compra y dar el vuelto.\n- Quiebres de stock por no tener un control automatizado de la mercancía que entra y sale.\n- Dificultad para registrar qué empleado realizó qué venta.')
    
    doc.add_heading('3. OBJETIVOS DEL SISTEMA', level=1)
    doc.add_heading('Objetivo General', level=2)
    doc.add_paragraph('Desarrollar e implementar un Sistema de Ventas de Escritorio (SysVentas) usando JavaFX y persistencia de datos local, que permita gestionar de manera integral el flujo comercial de la tienda.')
    doc.add_heading('Objetivos Específicos', level=2)
    doc.add_paragraph('- Automatizar el proceso de facturación y emisión de comprobantes.\n- Controlar en tiempo real el stock de productos, sus categorías y marcas.\n- Mantener un registro seguro de los clientes mediante su DNI/RUC.\n- Gestionar usuarios y perfiles (Administrador, Cajero) para controlar el acceso al sistema.')
    
    doc.add_heading('4. MÓDULOS DEL SISTEMA', level=1)
    doc.add_paragraph('El sistema SysVentas se compone de los siguientes módulos funcionales:\n\n1. Módulo de Seguridad (Login): Validación de usuarios registrados en la base de datos antes de permitir el acceso a las funciones del sistema.\n2. Módulo de Mantenimiento de Productos: Permite realizar operaciones CRUD sobre el catálogo de productos.\n3. Módulo de Clientes y Proveedores: Registro de los datos personales (DNI, Nombre, Dirección) para emitir los comprobantes correctamente.\n4. Módulo de Ventas y Comprobantes: Interfaz rápida donde el cajero puede agregar productos, ver el subtotal, calcular el total y finalizar la venta descontando el stock automáticamente.')

    doc.add_heading('5. ARQUITECTURA TECNOLÓGICA', level=1)
    doc.add_paragraph('Lenguaje: Java 17\nInterfaz Gráfica: JavaFX\nBase de Datos: H2 Database con HikariCP\nValidación: Hibernate Validator (Jakarta Validation)\nPatrón de Diseño: MVC y Arquitectura en Capas.')

    doc.save('01_Caso_de_Estudio_FINAL_COMPLETO.docx')
    print("Caso de estudio generado exitosamente.")

if __name__ == "__main__":
    crear_caso_estudio()
